"""Daily recap export: one row per issue, as CSV or Excel. Ready to paste into a tracker."""

from __future__ import annotations

import csv
import io
import json

from openpyxl import Workbook
from openpyxl.styles import Font

from .rules import ISSUE_ORDER, RANK, playbook_for

HEADERS = ["Date", "Priority", "Store", "Store name", "Area", "PC", "Issue", "Severity", "Reading", "Limit",
           "Why flagged", "Days in a row", "Trend", "New today", "Owner", "First step", "Handled note"]


def rows(con, cfg: dict, day: str) -> list[list]:
    stores = {r["code"]: dict(r) for r in con.execute("SELECT * FROM stores")}
    pcs = {r["pc_id"]: dict(r) for r in con.execute("SELECT * FROM pcs")}
    notes = {}
    for h in con.execute("SELECT * FROM handled WHERE date=? ORDER BY id", (day,)):
        notes[(h["store_code"], h["issue_type"])] = h["note"]
    items = [dict(r) for r in con.execute("SELECT * FROM issues WHERE date=?", (day,))]
    items.sort(key=lambda i: (-RANK[i["severity"]], ISSUE_ORDER.index(i["issue_type"]), -i["streak"], i["store_code"]))
    out = []
    for n, i in enumerate(items, 1):
        pb = playbook_for(cfg, i)
        out.append([day, n, i["store_code"], stores[i["store_code"]]["name"], stores[i["store_code"]]["area"],
                    pcs[i["pc_id"]]["role"] if i["pc_id"] else "whole store", cfg[i["issue_type"]]["label"],
                    i["severity"], i["value"], i["threshold"], i["reason"], i["streak"], i["trend"] or "",
                    "yes" if i["is_new"] else "", pb["owner"], pb["steps"][0],
                    notes.get((i["store_code"], i["issue_type"]), "")])
    return out


def to_csv(con, cfg, day) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    w.writerows(rows(con, cfg, day))
    return buf.getvalue()


def to_xlsx(con, cfg, day) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"Recap {day}"
    ws.append(HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows(con, cfg, day):
        ws.append(r)
    widths = [11, 8, 9, 20, 8, 12, 22, 9, 26, 22, 44, 8, 10, 8, 12, 60, 30]
    for col, w in zip("ABCDEFGHIJKLMNOPQ", widths):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    r = con.execute("SELECT text FROM resumes WHERE date=?", (day,)).fetchone()
    if r:
        ws2 = wb.create_sheet("Resume")
        for line in r["text"].splitlines():
            ws2.append([line])
        ws2.column_dimensions["A"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
