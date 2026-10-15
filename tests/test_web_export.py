import csv
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from pchealth import export
from pchealth.web import create_app
from tests.conftest import END


@pytest.fixture()
def client(filled, cfg):
    return TestClient(create_app(filled, cfg))


@pytest.mark.parametrize("path", ["/", "/stores", "/issues", "/resume", "/runs", "/healthz",
                                  "/stores?type=lan_degraded&show=all", "/stores?severity=high&owner=Vendor",
                                  "/?d=2026-09-20"])
def test_pages_render(client, path):
    r = client.get(path)
    assert r.status_code == 200


def test_store_detail_and_unknown(client, filled):
    code = filled.execute("SELECT store_code FROM issues WHERE date=? LIMIT 1", (END.isoformat(),)).fetchone()[0]
    r = client.get(f"/stores/{code}")
    assert r.status_code == 200 and "Last 30 days" in r.text and "<svg" in r.text
    assert client.get("/stores/NOPE").status_code == 404
    assert client.get("/?d=1999-01-01").status_code == 404


def test_filter_by_type(client, filled):
    r = client.get("/stores?type=edc_offline&show=all")
    codes = {x[0] for x in filled.execute("SELECT store_code FROM issues WHERE date=? AND issue_type='edc_offline'", (END.isoformat(),))}
    for c in codes:
        assert c in r.text


def test_mark_handled_and_undo(client, filled):
    row = filled.execute("SELECT store_code, issue_type FROM issues WHERE date=? LIMIT 1", (END.isoformat(),)).fetchone()
    r = client.post("/handled", data={"d": END.isoformat(), "store": row[0], "issue_type": row[1],
                                      "note": "Called the store", "back": f"/stores/{row[0]}"}, follow_redirects=False)
    assert r.status_code == 303
    page = client.get(f"/stores/{row[0]}").text
    assert "Called the store" in page
    hid = filled.execute("SELECT id FROM handled ORDER BY id DESC LIMIT 1").fetchone()[0]
    client.post(f"/handled/{hid}/undo", data={"back": "/"})
    assert "Called the store" not in client.get(f"/stores/{row[0]}").text
    assert client.post("/handled", data={"d": END.isoformat(), "store": row[0], "issue_type": "bogus"}).status_code == 400


def test_resume_page_has_copy_and_no_send(client):
    r = client.get("/resume")
    assert "Copy text" in r.text and "never sends" in r.text
    assert "send" not in r.text.lower().replace("never sends", "").replace("send it", "").replace("whether to send", "")


def test_exports(client, filled, cfg):
    r = client.get("/export.csv")
    assert r.status_code == 200
    rows = list(csv.reader(io.StringIO(r.text)))
    n = filled.execute("SELECT COUNT(*) FROM issues WHERE date=?", (END.isoformat(),)).fetchone()[0]
    assert rows[0] == export.HEADERS and len(rows) == n + 1
    x = client.get("/export.xlsx")
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[1] == "Resume" and wb.active.max_row == n + 1


def test_today_is_a_status_board(client, filled):
    from pchealth import fleet
    r = client.get("/")
    html = r.text
    # one tile per store, all 60, each with a status class
    n_stores = filled.execute("SELECT COUNT(*) FROM stores").fetchone()[0]
    assert html.count('class="tile st-') == n_stores
    for label in ("Healthy PCs", "Warning PCs", "Critical PCs", "Stores with no data", "vs yesterday"):
        assert label in html
    assert 'class="chart stack"' in html and 'class="mixbar"' in html
    # explanation text is not on the Today screen; it lives behind a click
    expl = filled.execute("SELECT explanations FROM resumes WHERE date=?", (END.isoformat(),)).fetchone()[0]
    import json
    first = next(iter(json.loads(expl).values()))
    assert first not in html


def test_today_filters_server_side(client, filled):
    code = filled.execute("SELECT store_code FROM issues WHERE date=? AND issue_type='edc_offline'", (END.isoformat(),)).fetchone()[0]
    html = client.get("/?status=healthy").text
    import re
    visible = re.findall(r'<li class="tile st-(\w+)[^"]*"[^>]*?(?<!hidden)>', html)
    assert visible and set(visible) == {"healthy"}
    html = client.get("/?type=edc_offline").text
    shown = re.findall(r'data-types="([^"]*)">', html)
    assert shown and all("edc_offline" in s for s in shown)
    assert code in html


def test_store_panel_fragment(client, filled):
    code = filled.execute("SELECT store_code FROM issues WHERE date=? AND pc_id IS NOT NULL LIMIT 1", (END.isoformat(),)).fetchone()[0]
    r = client.get(f"/stores/{code}/panel")
    assert r.status_code == 200
    assert "Reading" in r.text and "Limit" in r.text and "Owner" in r.text
    assert r.text.count('class="chart line"') >= 3
    assert "CPU temperature" in r.text and "Disk free" in r.text and "Days since restart" in r.text
    assert client.get("/stores/NOPE/panel").status_code == 404


def test_status_filter_on_store_list(client):
    assert client.get("/stores?status=critical&show=all").status_code == 200
